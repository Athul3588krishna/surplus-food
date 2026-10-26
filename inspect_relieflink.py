import docx
import xml.etree.ElementTree as ET

doc = docx.Document(r"C:\Users\User\Downloads\ReliefLink_Final_Report_v33.docx")
print("Total paragraphs in ReliefLink:", len(doc.paragraphs))
for i, p in enumerate(doc.paragraphs):
    xml_str = p._element.xml
    if "blip" in xml_str:
        root = ET.fromstring(xml_str)
        for elem in root.iter():
            if elem.tag.endswith("blip"):
                rId = elem.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                target = doc.part.rels[rId].target_ref
                txt = p.text.strip()
                prev_txt = doc.paragraphs[i-1].text.strip() if i > 0 else ""
                next_txt = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ""
                print(f"P{i}: target={target} | text={txt} | prev={prev_txt} | next={next_txt}")
