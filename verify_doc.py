import docx
import xml.etree.ElementTree as ET

doc = docx.Document(r"d:\surplus-food\Surplus_Food_Project_Report.docx")
print("Surplus report total paragraphs:", len(doc.paragraphs))
imgs = []
for i, p in enumerate(doc.paragraphs):
    xml_str = p._element.xml
    if "blip" in xml_str:
        root = ET.fromstring(xml_str)
        for elem in root.iter():
            if elem.tag.endswith("blip"):
                rId = elem.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                target = doc.part.rels[rId].target_ref
                prev_txt = doc.paragraphs[i-1].text.strip() if i > 0 else ""
                next_txt = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ""
                imgs.append((i, target, prev_txt, next_txt))

print(f"Total images in Surplus docx: {len(imgs)}")
for idx, target, prev, nxt in imgs:
    print(f"P{idx}: {target} | prev='{prev[:40]}' | next='{nxt[:40]}'")
